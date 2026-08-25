"""يولّد قالب Excel لإدخال المنشورات المرجعية المرتبطة بالمنتجات.

الورقة الأولى «المنشورات» تُملأ يدوياً، والورقة الثانية «المنتجات» مرجع
للقراءة فقط يبيّن رقم كل منتج (نفس ترقيم ملف الاستيراد).
"""
from __future__ import annotations

import sys

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HEADERS = [
    ("رقم المنتج", 12),
    ("نوع البوست", 20),
    ("هدف البوست", 26),
    ("الهوك", 34),
    ("الكابشن", 60),
    ("CTA", 26),
    ("الهاشتاغات", 34),
    ("وصف الصورة", 30),
    ("رابط الصورة", 30),
    ("ملاحظات", 22),
]

POST_TYPES = [
    "promotional",
    "highlight_product_and_offer",
    "educational_and_interactive",
    "brand_awareness",
    "testimonial",
    "seasonal",
]


def build(products: list[tuple[int, str, float | None, str | None]], out_path: str) -> None:
    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "المنشورات"
    ws.sheet_view.rightToLeft = True

    head_fill = PatternFill("solid", fgColor="1F4E79")
    head_font = Font(color="FFFFFF", bold=True, size=11)
    for i, (name, width) in enumerate(HEADERS, start=1):
        c = ws.cell(row=1, column=i, value=name)
        c.fill, c.font = head_fill, head_font
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 24

    # قائمة منسدلة لنوع البوست + تقييد رقم المنتج بالمدى الصحيح
    dv_type = DataValidation(
        type="list", formula1='"' + ",".join(POST_TYPES) + '"', allow_blank=True
    )
    ws.add_data_validation(dv_type)
    dv_type.add(f"B2:B1000")

    dv_ref = DataValidation(
        type="whole", operator="between", formula1=1, formula2=len(products),
        allow_blank=True, showErrorMessage=True,
        error=f"رقم المنتج يجب أن يكون بين 1 و {len(products)}",
        errorTitle="رقم غير صالح",
    )
    ws.add_data_validation(dv_ref)
    dv_ref.add("A2:A1000")

    for row in range(2, 1001):
        ws.cell(row=row, column=5).alignment = Alignment(wrap_text=True, vertical="top")

    # ورقة مرجعية للمنتجات
    ref = wb.create_sheet("المنتجات (مرجع)")
    ref.sheet_view.rightToLeft = True
    for i, (name, width) in enumerate(
        [("رقم المنتج", 12), ("اسم المنتج", 55), ("السعر", 12), ("الفئة", 24)], start=1
    ):
        c = ref.cell(row=1, column=i, value=name)
        c.fill, c.font = head_fill, head_font
        c.alignment = Alignment(horizontal="center")
        ref.column_dimensions[get_column_letter(i)].width = width
    for r, (ref_no, title, price, category) in enumerate(products, start=2):
        ref.cell(row=r, column=1, value=ref_no)
        ref.cell(row=r, column=2, value=title)
        ref.cell(row=r, column=3, value=price)
        ref.cell(row=r, column=4, value=category)
    ref.freeze_panes = "A2"

    wb.save(out_path)


def main() -> int:
    from database.models import Product, User
    from database.session import SessionLocal

    out_path = sys.argv[1] if len(sys.argv) > 1 else "alboraq_posts_template.xlsx"
    db = SessionLocal()
    try:
        admin = db.query(User).filter(User.role == "super_admin").first()
        rows = (
            db.query(Product.id, Product.title, Product.price, Product.category)
            .filter(Product.user_id == admin.id)
            .order_by(Product.id)
            .all()
        )
    finally:
        db.close()

    # رقم المنتج = ترتيبه (1..N) بنفس ترتيب ملف الاستيراد
    products = [(i, r[1], r[2], r[3]) for i, r in enumerate(rows, start=1)]
    build(products, out_path)
    print(f"تم إنشاء القالب: {out_path} ({len(products)} منتجاً في الورقة المرجعية)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
